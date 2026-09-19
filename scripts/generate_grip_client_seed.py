"""Generate the exact 560-unit client seed from the approved GRIP workbook."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/generated/GRIP_LIST_BARCODE_PAYLOADS.xlsx"
DESTINATION = ROOT / "lightbenders_warehouse/lightbenders_warehouse/setup/data/initial_inventory.json"
EXPECTED_ROWS = 560


def _slug(value: str) -> str:
	return re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")


def _product_codes(rows: list[dict[str, object]]) -> dict[tuple[str, str], str]:
	by_shortcut: defaultdict[str, set[str]] = defaultdict(set)
	for row in rows:
		by_shortcut[str(row["shortcut"])].add(str(row["equipment"]))

	codes: dict[tuple[str, str], str] = {}
	for shortcut, equipment_names in by_shortcut.items():
		for equipment in sorted(equipment_names):
			code = f"KUPO-{shortcut}"
			if len(equipment_names) > 1:
				code = f"{code}-{_slug(equipment)}"
			codes[(shortcut, equipment)] = code
	return codes


def read_rows() -> list[dict[str, object]]:
	if not SOURCE.exists():
		raise FileNotFoundError(f"Workbook not found: {SOURCE}")
	workbook = load_workbook(SOURCE, read_only=True, data_only=True)
	worksheet = workbook["Barcode Payloads"]
	headers = [str(value).strip() for value in next(worksheet.iter_rows(values_only=True))]
	required = {
		"Global No.",
		"Category",
		"Equipment",
		"Equipment Shortcut",
		"Unit No.",
		"Barcode Payload",
	}
	missing = required - set(headers)
	if missing:
		raise ValueError(f"Workbook is missing columns: {sorted(missing)}")

	rows: list[dict[str, object]] = []
	for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
		record = dict(zip(headers, values))
		if not all(record.get(column) not in (None, "") for column in required):
			raise ValueError(f"Workbook row {row_number} has a missing required value")
		rows.append(
			{
				"global_no": int(record["Global No."]),
				"category": str(record["Category"]).strip(),
				"equipment": str(record["Equipment"]).strip(),
				"shortcut": str(record["Equipment Shortcut"]).strip(),
				"unit_no": int(record["Unit No."]),
				"barcode": str(record["Barcode Payload"]).strip(),
			}
		)
	return rows


def validate(rows: list[dict[str, object]]) -> None:
	if len(rows) != EXPECTED_ROWS:
		raise ValueError(f"Expected {EXPECTED_ROWS} rows, found {len(rows)}")
	barcodes = [str(row["barcode"]) for row in rows]
	if len(set(barcodes)) != len(barcodes):
		raise ValueError("Barcode payloads are not unique")
	global_numbers = [int(row["global_no"]) for row in rows]
	if sorted(global_numbers) != list(range(1, EXPECTED_ROWS + 1)):
		raise ValueError("Global No. values must be exactly 1 through 560")


def generate() -> None:
	rows = read_rows()
	validate(rows)
	codes = _product_codes(rows)
	result = [
		{
			"product_code": codes[(str(row["shortcut"]), str(row["equipment"]))],
			"equipment_name": row["equipment"],
			"category": row["category"],
			"tracking_mode": "SERIALIZED",
			"quantity": 1,
			"asset_id": f"GRIP-{int(row['global_no']):04d}",
			"barcode_payload": row["barcode"],
		}
		for row in rows
	]
	DESTINATION.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
	print(f"Generated {len(result)} serialized units and {len(codes)} items")


if __name__ == "__main__":
	generate()
